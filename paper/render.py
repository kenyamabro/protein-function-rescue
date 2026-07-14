"""Render the manuscript to DOCX (and PDF if a converter is available).

Usage:  python paper/render.py
Requires: pypandoc-binary (bundles pandoc). PDF additionally needs either
Microsoft Word (via docx2pdf) or a LaTeX engine.
"""
from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
MD = HERE / "protein_function_rescue.md"
DOCX = HERE / "protein_function_rescue.docx"
PDF = HERE / "protein_function_rescue.pdf"
TITLE = ("Structure-based functional annotation of hypothetical proteins "
         "using the AlphaFold Database and TM-align")


def format_for_bmc():
    """Apply BMC Bioinformatics submission formatting to the Pandoc DOCX."""
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document(DOCX)
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        if style_name in doc.styles:
            doc.styles[style_name].paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        sect_pr = section._sectPr
        for old in sect_pr.findall(qn("w:lnNumType")):
            sect_pr.remove(old)
        line_numbers = OxmlElement("w:lnNumType")
        line_numbers.set(qn("w:countBy"), "1")
        line_numbers.set(qn("w:restart"), "continuous")
        sect_pr.append(line_numbers)

        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer_p.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for node in (begin, instr, separate, text, end):
            run._r.append(node)

    doc.save(DOCX)


def render_docx():
    import pypandoc
    import os
    cwd = os.getcwd()
    os.chdir(HERE)  # so relative image paths (figures/, ../benchmark/) resolve
    try:
        pypandoc.convert_file(
            MD.name, "docx", outputfile=DOCX.name,
            extra_args=["--resource-path=.:.."],
        )
        format_for_bmc()
    finally:
        os.chdir(cwd)
    print(f"wrote {DOCX}")


def render_pdf():
    try:
        from docx2pdf import convert
        convert(str(DOCX), str(PDF))
        if PDF.exists():
            print(f"wrote {PDF}")
            return True
    except Exception as exc:
        print(f"PDF via Word failed ({exc})")
    # Fallback: pandoc with a LaTeX engine, if present.
    try:
        import shutil, pypandoc
        for engine in ("xelatex", "pdflatex", "tectonic"):
            if shutil.which(engine):
                pypandoc.convert_file(str(MD), "pdf", outputfile=str(PDF),
                                      extra_args=[f"--pdf-engine={engine}",
                                                  "--resource-path=" + str(HERE) + ":" + str(HERE.parent)])
                print(f"wrote {PDF} (via {engine})")
                return True
    except Exception as exc:
        print(f"PDF via LaTeX failed ({exc})")
    print("PDF not produced (install MS Word or a LaTeX engine); DOCX is available.")
    return False


if __name__ == "__main__":
    render_docx()
    render_pdf()
